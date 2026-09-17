# -*- coding: utf-8 -*-
r"""
A AGENDA É A FONTE — leitura SOMENTE LEITURA de `Agenda_Licitacoes.xlsx`.

Cópia enxuta de `modulo2/utils.ler_linhas_agenda()` e de
`modulo2/utils.mapear_colunas_agenda()`, pela mesma razão de
`convencao_edital.py`: a regra é provada, roda em produção, e não vale um import
cruzado entre repositórios. Vale a mesma advertência — a suíte de equivalência
roda esta leitura e a do Módulo 2 sobre a MESMA planilha e exige linhas
idênticas.

AS COLUNAS SÃO LOCALIZADAS PELO TEXTO DO CABEÇALHO, NUNCA POR POSIÇÃO
    É a regra do Módulo 2, e o motivo está registrado lá: com índices fixos,
    bastava o Módulo 1 inserir uma coluna no meio da Agenda para a leitura passar
    a pegar a coluna errada — em silêncio, com dados plausíveis no lugar certo.
    Ou o cabeçalho é encontrado, ou a leitura falha explicitamente.

    A medição desta rodada encontrou 52 colunas na aba "2026", e "Código
    ConLicitação" na coluna 30 (AD). Nenhum desses números está no código: os
    dois são resolvidos em tempo de execução, e é por isso que podem mudar sem
    quebrar nada.

A AGENDA NUNCA É ABERTA PARA ESCRITA
    `read_only=True`, e não há — nem pode haver — um `wb.save()` neste arquivo.
    A Agenda é o banco de dados oficial do projeto, escrito exclusivamente pelo
    Módulo 1.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from convencao_edital import remover_acentos, texto_codigo

# Mesma convenção do Módulo 1 e do Módulo 2: cabeçalho na linha 2, dados a
# partir da linha 3.
LINHA_CABECALHO_AGENDA = 2
PRIMEIRA_LINHA_DADOS = LINHA_CABECALHO_AGENDA + 1

# Só o que o Inventariador precisa para IDENTIFICAR e LOCALIZAR um edital.
# Deliberadamente menor que o CABECALHOS_LIDOS do Módulo 2: o Inventariador não
# classifica nada, então objeto, categoria, tipo, plataforma e afins não têm
# papel aqui. O que ele não usa, ele não lê.
CABECALHOS_LIDOS = {
    "cidade": "CIDADE",
    "uf": "UF",
    "numero_pregao": "Número do Pregão",
    "codigo_conlicitacao": "Código ConLicitação",
    "data_limite": "Data Limite para Cadastro Proposta",
}


class CabecalhoAgendaAusente(RuntimeError):
    """
    Um cabeçalho esperado não existe na planilha.

    Mesma classe de erro do Módulo 2, e pelo mesmo motivo: é melhor parar
    dizendo qual cabeçalho faltou do que devolver linhas montadas com a coluna
    errada.
    """


def _normalizar_cabecalho(valor) -> str:
    """
    Forma canônica de um cabeçalho para comparação: sem acento, maiúsculo e com
    espaços colapsados.

    Cópia de utils._normalizar_cabecalho(). Tolera diferença de
    acentuação/caixa/espaço (alguém reescreve "Nº DA LINHA " à mão), mas não
    tolera coluna ausente. Copiar isto ERRADO — comparando o texto cru — faria a
    leitura falhar num cabeçalho que o Módulo 2 aceita, e a suíte de
    equivalência é quem cobra que as duas normalizações sejam a mesma.
    """
    if valor is None:
        return ""
    return " ".join(remover_acentos(str(valor)).upper().split())


def mapear_colunas_agenda(celulas_cabecalho) -> dict:
    """
    {campo: índice 0-based} a partir da LINHA DE CABEÇALHO real da planilha.

    Cópia de utils.mapear_colunas_agenda(). Nunca devolve mapa parcial: ou acha
    todos os cabeçalhos de CABECALHOS_LIDOS, ou levanta CabecalhoAgendaAusente
    listando os que faltaram.
    """
    posicoes = {}
    for indice, celula in enumerate(celulas_cabecalho):
        titulo = _normalizar_cabecalho(getattr(celula, "value", celula))
        if titulo and titulo not in posicoes:  # primeira ocorrência vence
            posicoes[titulo] = indice

    mapa, faltando = {}, []
    for campo, cabecalho in CABECALHOS_LIDOS.items():
        indice = posicoes.get(_normalizar_cabecalho(cabecalho))
        if indice is None:
            faltando.append(cabecalho)
        else:
            mapa[campo] = indice

    if faltando:
        raise CabecalhoAgendaAusente(
            "Agenda: cabeçalho(s) não encontrado(s) na linha "
            f"{LINHA_CABECALHO_AGENDA}: "
            + ", ".join(f"'{c}'" for c in faltando)
            + ". Cabeçalhos presentes: "
            + ", ".join(f"'{t}'" for t in posicoes)
            + ". O Inventariador NÃO lê a Agenda por posição de coluna."
        )
    return mapa


def ler_linhas_agenda(caminho_agenda: Path, aba: str) -> list[dict]:
    """
    Uma linha por edital, já com os campos nomeados.

    Linhas sem Cidade ou sem Número do Pregão são ignoradas — é a mesma regra do
    Módulo 2, e o motivo é o mesmo: sem essas duas informações não dá para
    descrever o edital nem montar o nome da pasta.

    A ordem devolvida é a ORDEM DA PLANILHA, e `linha_agenda` carrega o número
    real da linha. Quem ordena o inventário é o inventário (§4.4); aqui a ordem é
    a da fonte, para a comparação com o Módulo 2 ser linha a linha.

    UMA DIVERGÊNCIA DELIBERADA EM RELAÇÃO AO MÓDULO 2, E ESTÁ DECLARADA
        O Módulo 2 faz `(valor("cidade") or "").strip()` — se a célula vier como
        número (o Excel converteu), aquilo levanta AttributeError e a execução
        inteira para. Aqui a célula passa por `str()` antes do `.strip()`.

        A diferença só aparece em célula NÃO textual, e a escolha segue o §4.7:
        um edital estranho não pode derrubar os outros. Para toda célula de
        texto — que é o caso das 2040 linhas medidas nesta rodada — as duas
        leituras devolvem exatamente a mesma coisa, e é isso que a suíte de
        equivalência confere linha a linha.
    """
    caminho_agenda = Path(caminho_agenda)
    if not caminho_agenda.exists():
        raise FileNotFoundError(f"Agenda não encontrada em '{caminho_agenda}'")

    wb = openpyxl.load_workbook(caminho_agenda, read_only=True, data_only=True)
    try:
        if aba not in wb.sheetnames:
            raise CabecalhoAgendaAusente(
                f"Aba {aba!r} não encontrada na Agenda '{caminho_agenda}'. "
                f"Abas presentes: {wb.sheetnames}"
            )
        ws = wb[aba]

        linhas_planilha = ws.iter_rows(min_row=LINHA_CABECALHO_AGENDA)
        celulas_cabecalho = next(linhas_planilha, None)
        if celulas_cabecalho is None:
            raise CabecalhoAgendaAusente(
                f"Agenda '{caminho_agenda}', aba {aba!r}: a linha "
                f"{LINHA_CABECALHO_AGENDA} (cabeçalho) não existe."
            )
        colunas = mapear_colunas_agenda(celulas_cabecalho)

        linhas = []
        for numero_linha, linha in enumerate(linhas_planilha, start=PRIMEIRA_LINHA_DADOS):
            def valor(campo):
                indice = colunas[campo]
                return linha[indice].value if indice < len(linha) else None

            cidade = (valor("cidade") or "")
            uf = (valor("uf") or "")
            num_pregao = (valor("numero_pregao") or "")
            cidade = str(cidade).strip()
            uf = str(uf).strip()
            num_pregao = str(num_pregao).strip()
            if not cidade or not num_pregao:
                continue

            linhas.append({
                "linha_agenda": numero_linha,
                "cidade": cidade,
                "uf": uf,
                "numero_pregao": num_pregao,
                "codigo_conlicitacao": texto_codigo(valor("codigo_conlicitacao")),
                "data_limite": _texto_simples(valor("data_limite")),
            })
        return linhas
    finally:
        wb.close()


def _texto_simples(valor) -> str:
    """
    A data-limite entra no inventário só como TEXTO, do jeito que está na
    célula. O Inventariador não interpreta data, não calcula prazo e não decide
    nada com ela — ela viaja para quem for consumir o inventário conseguir
    reconhecer a linha sem reabrir a Agenda.
    """
    if valor is None:
        return ""
    return str(valor).strip()
